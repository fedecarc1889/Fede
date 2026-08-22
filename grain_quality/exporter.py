"""Exportación de resultados a Excel (SDD sección 20).

Genera un libro con cuatro hojas:
    - "Resumen Contratos":  una fila por contrato.
    - "Resultados CTG":     una fila por CTG.
    - "Detalle Calidad":    una fila por contrato + CTG + rubro.
    - "Errores":            filas rechazadas y motivo.
"""

from __future__ import annotations

from pathlib import Path
from typing import IO, Union

from openpyxl import Workbook
from openpyxl.styles import Font

from .models import ContractResult, CTGResult, ImportError_

_HEADER_FONT = Font(bold=True)


def _write_headers(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT


def export_results(
    contract_results: list[ContractResult],
    errors: list[ImportError_],
    target: Union[str, Path, IO[bytes]],
) -> None:
    workbook = Workbook()

    # Hoja "Resumen Contratos"
    sheet = workbook.active
    sheet.title = "Resumen Contratos"
    _write_headers(sheet, [
        "Contrato", "Cant. CTG", "Toneladas", "Consolidación",
        "Bonificación %", "Rebaja %", "Neto %", "Merma %", "Advertencias",
    ])
    for contract in contract_results:
        sheet.append([
            contract.contract,
            len(contract.ctgs),
            contract.total_weight,
            "Ponderada por toneladas" if contract.weighted else "Promedio simple",
            round(contract.total_bonificacion, 4),
            round(contract.total_rebaja, 4),
            round(contract.neto, 4),
            round(contract.total_merma, 4),
            "; ".join(contract.warnings),
        ])

    # Hoja "Resultados CTG"
    sheet = workbook.create_sheet("Resultados CTG")
    _write_headers(sheet, [
        "Contrato", "CTG", "Producto", "Toneladas", "Norma", "Versión",
        "Bonificación %", "Rebaja %", "Neto %", "Merma %",
        "Fecha cálculo", "Advertencias",
    ])
    all_ctgs: list[CTGResult] = [
        ctg for contract in contract_results for ctg in contract.ctgs
    ]
    for ctg in all_ctgs:
        sheet.append([
            ctg.contract, ctg.ctg, ctg.product_name, ctg.weight_tn,
            ctg.norm_reference, ctg.norm_version,
            round(ctg.total_bonificacion, 4),
            round(ctg.total_rebaja, 4),
            round(ctg.neto, 4),
            round(ctg.total_merma, 4),
            ctg.calculated_at.strftime("%Y-%m-%d %H:%M:%S"),
            "; ".join(ctg.warnings),
        ])

    # Hoja "Detalle Calidad"
    sheet = workbook.create_sheet("Detalle Calidad")
    _write_headers(sheet, [
        "Contrato", "CTG", "Producto", "Rubro", "Unidad",
        "Planta", "Cámara", "Valor válido", "Origen",
        "Base", "Tolerancia", "Norma", "Versión", "Regla",
        "Tipo ajuste", "Ajuste %",
    ])
    for ctg in all_ctgs:
        for rubro in ctg.rubros:
            sheet.append([
                ctg.contract, ctg.ctg, ctg.product_name,
                rubro.rubro_name, rubro.unit,
                rubro.plant_value, rubro.chamber_value,
                rubro.selected_value,
                rubro.selected_source.value if rubro.selected_source else "",
                rubro.base, rubro.tolerance,
                rubro.norm_reference, rubro.norm_version,
                rubro.rule_description,
                rubro.adjustment_type.value,
                rubro.signed_percentage,
            ])

    # Hoja "Errores"
    sheet = workbook.create_sheet("Errores")
    _write_headers(sheet, ["Fila", "Contrato", "CTG", "Motivo del rechazo"])
    for error in errors:
        sheet.append([error.row_number, error.contract, error.ctg, error.reason])

    workbook.save(target)
