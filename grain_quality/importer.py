"""Importador de archivos Excel (SDD secciones 9, 10 y 16).

Formato esperado (primera hoja del .xlsx, encabezados en la fila 1):

    Contrato | CTG | Producto | Toneladas | <Rubro> Planta | <Rubro> Cámara | ...

Los encabezados de rubro se detectan de forma flexible: se normalizan
(mayúsculas, sin acentos) y se acepta tanto el código interno
(``materias_extranas``) como el nombre visible (``Materias extrañas``),
seguido de ``PLANTA`` o ``CAMARA``. También se aceptan abreviaturas
configuradas en ``_ALIASES``.

Los errores de una fila no impiden procesar el resto del archivo: las
filas rechazadas se informan con número de fila y motivo.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import IO, Optional, Union

from openpyxl import load_workbook

from .calculator import calculate_ctg, consolidate_contracts
from .models import ContractResult, CTGResult, ImportError_, MeasurementInput
from .norms_registry import NormNotFoundError, NormsRegistry, _normalize, default_registry

# Abreviaturas habituales en planillas comerciales -> código de rubro interno
_ALIASES = {
    "ME": "MATERIAS_EXTRANAS",
    "MMEE": "MATERIAS_EXTRANAS",
    "CE": "CUERPOS_EXTRANOS",
    "HUM": "HUMEDAD",
    "DANADOS": "GRANOS_DANADOS",
    "VERDES": "GRANOS_VERDES",
    "QUEBRADOS": "GRANOS_QUEBRADOS",
    "MG": "MATERIA_GRASA",
    "PH": "PESO_HECTOLITRICO",
}

_CONTRACT_HEADERS = {"CONTRATO", "CONTRACT", "NRO CONTRATO", "NUMERO CONTRATO"}
_CTG_HEADERS = {"CTG", "NRO CTG", "NUMERO CTG"}
_PRODUCT_HEADERS = {"PRODUCTO", "PRODUCT", "GRANO", "CEREAL"}
_WEIGHT_HEADERS = {"TONELADAS", "TN", "TONS", "PESO", "KILOS", "KG", "KILOGRAMOS"}
_SOURCE_SUFFIXES = {"PLANTA": "PLANTA", "CAMARA": "CAMARA"}


@dataclass
class ImportResult:
    """Resultado completo de una importación."""

    ctg_results: list[CTGResult] = field(default_factory=list)
    contract_results: list[ContractResult] = field(default_factory=list)
    errors: list[ImportError_] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def parse_number(raw) -> Optional[float]:
    """Convierte una celda a número. Vacío -> None (nunca cero).

    Acepta coma o punto decimal ("14,2" y "14.2"). Lanza ValueError si el
    contenido no es numérico.
    """
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = str(raw).strip()
    if text == "":
        return None
    text = text.replace(" ", "").replace("%", "")
    if "," in text and "." in text:
        # formato "1.234,56": el punto es separador de miles
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    return float(text)


def _canonical_rubro(name: str) -> str:
    norm = _normalize(name).replace(" ", "_").replace("/", "_")
    return _ALIASES.get(norm.replace("_", " ").strip(), _ALIASES.get(norm, norm))


def _map_headers(headers: list[Optional[str]]) -> tuple[dict, list[tuple[int, str, str]], list[str]]:
    """Clasifica encabezados en columnas fijas y columnas de rubro.

    Devuelve (columnas_fijas, columnas_rubro, advertencias) donde
    columnas_rubro es una lista de (índice, código_rubro, origen).
    """
    fixed: dict[str, int] = {}
    rubro_cols: list[tuple[int, str, str]] = []
    warnings: list[str] = []

    for idx, raw in enumerate(headers):
        if raw is None or str(raw).strip() == "":
            continue
        header = _normalize(str(raw))
        if header in _CONTRACT_HEADERS:
            fixed["contract"] = idx
        elif header in _CTG_HEADERS:
            fixed["ctg"] = idx
        elif header in _PRODUCT_HEADERS:
            fixed["product"] = idx
        elif header in _WEIGHT_HEADERS:
            fixed["weight"] = idx
            fixed["weight_unit"] = header  # para convertir kilos a toneladas
        else:
            matched = False
            for suffix, source in _SOURCE_SUFFIXES.items():
                if header.endswith(" " + suffix):
                    rubro = _canonical_rubro(header[: -len(suffix)].strip())
                    rubro_cols.append((idx, rubro, source))
                    matched = True
                    break
            if not matched:
                warnings.append(f"Columna no reconocida (ignorada): '{raw}'")
    return fixed, rubro_cols, warnings


def import_excel(
    file: Union[str, Path, IO[bytes]],
    registry: Optional[NormsRegistry] = None,
) -> ImportResult:
    """Ejecuta el proceso completo de importación (SDD sección 10)."""
    registry = registry or default_registry()
    result = ImportResult()

    # Paso 1 – Lectura
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # archivo corrupto o formato inválido
        result.errors.append(
            ImportError_(None, "", "", f"No se pudo leer el archivo Excel: {exc}")
        )
        return result

    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = list(next(rows))
    except StopIteration:
        result.errors.append(ImportError_(None, "", "", "El archivo está vacío."))
        return result

    # Paso 2 – Validación estructural de columnas
    fixed, rubro_cols, header_warnings = _map_headers(headers)
    result.warnings.extend(header_warnings)

    missing = [
        label
        for key, label in (("ctg", "CTG"), ("product", "Producto"))
        if key not in fixed
    ]
    if missing:
        result.errors.append(
            ImportError_(
                None, "", "",
                f"Faltan columnas obligatorias: {', '.join(missing)}."
            )
        )
        return result
    if not rubro_cols:
        result.errors.append(
            ImportError_(
                None, "", "",
                "No se detectó ninguna columna de rubro "
                "('<Rubro> Planta' / '<Rubro> Cámara')."
            )
        )
        return result

    seen_ctgs: set[tuple[str, str]] = set()

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(idx: Optional[int]):
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        contract = str(cell(fixed.get("contract")) or "").strip()
        ctg = str(cell(fixed.get("ctg")) or "").strip()
        product = str(cell(fixed.get("product")) or "").strip()

        # Paso 2 (cont.) – validaciones por fila
        if not ctg:
            result.errors.append(
                ImportError_(row_number, contract, ctg, "CTG vacío.")
            )
            continue
        if not product:
            result.errors.append(
                ImportError_(row_number, contract, ctg, "Producto vacío.")
            )
            continue
        if not registry.has_product(product):
            result.errors.append(
                ImportError_(
                    row_number, contract, ctg,
                    f"Producto inexistente o sin norma configurada: '{product}'."
                )
            )
            continue
        key = (contract, ctg)
        if key in seen_ctgs:
            result.errors.append(
                ImportError_(
                    row_number, contract, ctg,
                    f"CTG duplicado en el archivo (contrato '{contract}')."
                )
            )
            continue
        seen_ctgs.add(key)

        weight_tn: Optional[float] = None
        try:
            weight_tn = parse_number(cell(fixed.get("weight")))
            if weight_tn is not None and fixed.get("weight_unit") in {
                "KILOS", "KG", "KILOGRAMOS",
            }:
                weight_tn /= 1000.0
        except ValueError:
            result.errors.append(
                ImportError_(
                    row_number, contract, ctg,
                    f"Valor de peso no numérico: '{cell(fixed.get('weight'))}'."
                )
            )
            continue

        # Paso 3 – Normalización al modelo interno
        measurements: dict[str, MeasurementInput] = {}
        row_error = None
        for idx, rubro, source in rubro_cols:
            try:
                value = parse_number(cell(idx))
            except ValueError:
                row_error = (
                    f"Valor no numérico '{cell(idx)}' en columna '{headers[idx]}'."
                )
                break
            if value is None:
                continue
            m = measurements.setdefault(rubro, MeasurementInput(rubro=rubro))
            if source == "PLANTA":
                m.plant_value = value
            else:
                m.chamber_value = value
        if row_error:
            result.errors.append(ImportError_(row_number, contract, ctg, row_error))
            continue

        # Pasos 4 a 6 – valor válido, norma y cálculo (mismo motor que la
        # carga manual, CA-006/CA-007)
        try:
            ctg_result = calculate_ctg(
                product_code=product,
                measurements=list(measurements.values()),
                contract=contract,
                ctg=ctg,
                weight_tn=weight_tn,
                registry=registry,
            )
        except NormNotFoundError as exc:
            result.errors.append(ImportError_(row_number, contract, ctg, str(exc)))
            continue
        result.ctg_results.append(ctg_result)

    # Paso 7 – Consolidación por contrato
    result.contract_results = consolidate_contracts(result.ctg_results)
    return result
