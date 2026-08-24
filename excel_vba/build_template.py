"""Genera el libro CalculadoraCalidadGranos.xlsx a partir de las normas
JSON de grain_quality/norms/ (única fuente de verdad de las normas).

El libro contiene las hojas de configuración parametrizadas, la carga
manual, la hoja de datos para importación masiva, las hojas de resultados
y la trazabilidad. Los módulos VBA de excel_vba/src/ se importan luego en
el editor VBA (ver INSTRUCCIONES.md) y el libro se guarda como .xlsm.

Uso: python excel_vba/build_template.py [salida.xlsx]
"""

from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
NORMS_DIR = ROOT / "grain_quality" / "norms"

ARIAL = "Arial"
VERDE = "2E7D32"
AMARILLO = "FFFFCC"

HEADER_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", fgColor=VERDE)
BASE_FONT = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, bold=True, size=10)
INPUT_FILL = PatternFill("solid", fgColor=AMARILLO)

TIPO_MAP = {
    "rebaja_por_exceso": "REBAJA_POR_EXCESO",
    "merma_por_exceso": "MERMA_POR_EXCESO",
    "bonificacion_rebaja_lineal": "BONIFICACION_REBAJA_LINEAL",
    "escala": "ESCALA",
    "sin_ajuste": "SIN_AJUSTE",
}


def style_row(ws, row, n_cols, header=False):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT if header else BASE_FONT
        if header:
            cell.fill = HEADER_FILL


def write_headers(ws, headers, widths=None):
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_row(ws, 1, len(headers), header=True)
    for c, w in enumerate(widths or [14] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def load_norms():
    productos, normas, escalas = [], [], []
    for path in sorted(NORMS_DIR.glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        prod = data["product"]
        productos.append((prod["code"].upper(), prod["name"], "SI"))
        for version in data["versions"]:
            for rubro in version["rubros"]:
                calc = rubro["calc"]
                tipo = TIPO_MAP[calc["type"]]
                normas.append([
                    prod["code"].upper(),
                    str(version["version"]),
                    version["norm_reference"],
                    date.fromisoformat(version["valid_from"]),
                    date.fromisoformat(version["valid_to"]) if version.get("valid_to") else None,
                    version.get("status", "Activa"),
                    rubro["code"].upper(),
                    rubro["name"],
                    rubro["unit"],
                    "SI" if rubro.get("required") else "NO",
                    tipo,
                    calc.get("base"),
                    calc.get("tolerance", 0),
                    calc.get("factor", 0),
                    calc.get("bonif_factor", 0),
                    calc.get("rebaja_factor", 0),
                    calc.get("manipuleo", 0),
                    calc.get("max_pct"),
                    rubro.get("min"),
                    rubro.get("max"),
                    calc.get("formula", ""),
                ])
                for rng in calc.get("ranges", []):
                    escalas.append([
                        prod["code"].upper(), str(version["version"]),
                        rubro["code"].upper(), rng.get("from"), rng.get("to"),
                        rng.get("tipo", "SIN_AJUSTE"), rng.get("pct", 0),
                    ])
    return productos, normas, escalas


def build(target: Path) -> None:
    productos, normas, escalas = load_norms()
    wb = Workbook()

    # ---------- Inicio ----------
    ws = wb.active
    ws.title = "Inicio"
    ws.column_dimensions["B"].width = 110
    lineas = [
        ("CALCULADORA DE CALIDAD DE MERCADERÍA DE GRANOS", True),
        ("Motor parametrizable de liquidación de calidad (bonificaciones, rebajas y mermas).", False),
        ("", False),
        ("PUESTA EN MARCHA (una sola vez):", True),
        ("1. Habilite las macros y abra el editor VBA (Alt+F11).", False),
        ("2. Importe los módulos de la carpeta excel_vba/src (Archivo → Importar archivo…): "
         "todos los .bas y .cls.", False),
        ("3. Guarde el libro como 'Libro de Excel habilitado para macros (*.xlsm)'.", False),
        ("4. Ejecute la macro 'Inicializar' (Alt+F8) para crear los botones.", False),
        ("", False),
        ("HOJAS:", True),
        ("• CargaManual: cálculo individual. Elija producto, presione 'Cargar rubros', "
         "ingrese valores Planta/Cámara (celdas amarillas) y presione 'Calcular'.", False),
        ("• Datos: importación masiva. Pegue filas o use 'Importar archivo…' y luego "
         "'Procesar datos'. Los resultados van a Resumen Contratos / Resultados CTG / "
         "Detalle Calidad / Errores; 'Exportar resultados' los guarda en un libro nuevo.", False),
        ("• Config_Productos / Config_Normas / Config_Escalas: normas parametrizadas. "
         "Un cambio de norma se registra agregando filas de una NUEVA versión con su "
         "vigencia; las versiones históricas no se sobrescriben.", False),
        ("• Trazabilidad: registro de cada cálculo (valores, origen, norma, versión, "
         "regla, resultado, fecha y hora).", False),
        ("", False),
        ("REGLA CENTRAL (evaluada rubro por rubro): si un rubro tiene valor de Planta y "
         "de Cámara se usa el de CÁMARA; si tiene uno solo, se usa ese; si no tiene "
         "ninguno el valor es NULO (una celda vacía NUNCA es cero; el 0 sí es un valor "
         "válido).", False),
        ("", False),
        ("NOTA: los valores de bases y factores cargados en Config_Normas son "
         "configuración de referencia; valídelos contra la norma vigente de cada "
         "producto antes de usarlos comercialmente.", False),
        ("", False),
        ("La hoja Datos incluye filas de ejemplo (algunas con errores a propósito, para "
         "mostrar la hoja Errores). Bórrelas antes de cargar datos reales.", False),
    ]
    for i, (texto, bold) in enumerate(lineas, start=2):
        cell = ws.cell(row=i, column=2, value=texto)
        cell.font = Font(name=ARIAL, bold=bold, size=12 if i == 2 else 10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------- Config_Productos ----------
    ws = wb.create_sheet("Config_Productos")
    write_headers(ws, ["Codigo", "Nombre", "Activo"], [14, 22, 10])
    for r, fila in enumerate(productos, start=2):
        for c, v in enumerate(fila, start=1):
            ws.cell(row=r, column=c, value=v).font = BASE_FONT

    # ---------- Config_Normas ----------
    ws = wb.create_sheet("Config_Normas")
    write_headers(ws, [
        "Producto", "Version", "NormaReferencia", "VigenciaDesde", "VigenciaHasta",
        "Estado", "RubroCodigo", "RubroNombre", "Unidad", "Obligatorio",
        "TipoCalculo", "Base", "Tolerancia", "Factor", "FactorBonif",
        "FactorRebaja", "Manipuleo", "MaxPct", "Min", "Max", "Formula",
    ], [10, 8, 42, 13, 13, 9, 22, 26, 8, 11, 26, 7, 10, 7, 11, 12, 10, 8, 6, 6, 55])
    for r, fila in enumerate(normas, start=2):
        for c, v in enumerate(fila, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BASE_FONT
            if c in (4, 5) and v is not None:
                cell.number_format = "dd/mm/yyyy"

    # ---------- Config_Escalas ----------
    ws = wb.create_sheet("Config_Escalas")
    write_headers(ws, ["Producto", "Version", "Rubro", "Desde", "Hasta", "Tipo", "Pct"],
                  [10, 8, 20, 8, 8, 14, 7])
    for r, fila in enumerate(escalas, start=2):
        for c, v in enumerate(fila, start=1):
            ws.cell(row=r, column=c, value=v).font = BASE_FONT

    # ---------- CargaManual ----------
    ws = wb.create_sheet("CargaManual")
    for col, w in zip("ABCDEFGHIJ", [2, 16, 18, 8, 13, 13, 12, 12, 14, 10]):
        ws.column_dimensions[col].width = w
    campos = ["Producto", "Contrato", "CTG", "Toneladas", "Fecha análisis", "Observaciones"]
    for i, campo in enumerate(campos, start=2):
        ws.cell(row=i, column=2, value=campo).font = BOLD
        ws.cell(row=i, column=3).fill = INPUT_FILL
        ws.cell(row=i, column=3).font = BASE_FONT
    ws["C2"] = "SOJA"
    dv = DataValidation(type="list", formula1="=Config_Productos!$A$2:$A$100",
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(ws["C2"])
    ws.cell(row=8, column=2, value="(la norma aplicable aparece aquí al cargar los rubros)"
            ).font = Font(name=ARIAL, italic=True, size=9)

    encabezados = ["Rubro", "Unidad", "Base", "Valor Planta", "Valor Cámara",
                   "Valor válido", "Origen", "Tipo ajuste", "Ajuste %"]
    for c, h in enumerate(encabezados, start=2):
        cell = ws.cell(row=9, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    # Rubros de SOJA precargados con una fila de ejemplo de valores
    soja = [n for n in normas if n[0] == "SOJA"]
    ejemplo = {"HUMEDAD": (14.2, 13.8), "MATERIAS_EXTRANAS": (1.5, None)}
    for i, n in enumerate(soja):
        fila = 10 + i
        nombre = n[7] + (" *" if n[9] == "SI" else "")
        ws.cell(row=fila, column=2, value=nombre).font = BASE_FONT
        ws.cell(row=fila, column=3, value=n[8]).font = BASE_FONT
        ws.cell(row=fila, column=4, value=n[11]).font = BASE_FONT
        planta, camara = ejemplo.get(n[6], (None, None))
        cp = ws.cell(row=fila, column=5, value=planta)
        cc = ws.cell(row=fila, column=6, value=camara)
        cp.fill = INPUT_FILL
        cc.fill = INPUT_FILL
        cp.font = BASE_FONT
        cc.font = BASE_FONT
    nota = ("* rubro obligatorio · Celdas amarillas = carga de valores. Si un rubro tiene "
            "Planta y Cámara se usa CÁMARA. Vacío = sin análisis (no es cero). Los valores "
            "de ejemplo (Humedad 14,2/13,8 y Materias extrañas 1,5) son ilustrativos: "
            "reemplácelos por los del análisis real.")
    cell = ws.cell(row=10 + len(soja) + 1, column=2, value=nota)
    cell.font = Font(name=ARIAL, italic=True, size=9)
    cell.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=10 + len(soja) + 1, start_column=2,
                   end_row=10 + len(soja) + 2, end_column=10)

    # ---------- Datos (importación) ----------
    ws = wb.create_sheet("Datos")
    write_headers(ws, [
        "Contrato", "CTG", "Producto", "Toneladas",
        "Humedad Planta", "Humedad Camara",
        "Materias extranas Planta", "Materias extranas Camara",
        "Granos danados Planta", "Granos danados Camara",
        "Materia grasa Planta", "Materia grasa Camara",
    ], [10, 8, 10, 10, 13, 13, 20, 20, 18, 18, 17, 17])
    ejemplos = [
        ["10001", "CTG01", "SOJA", 30, 14.2, 13.8, 1.5, 1.1, 2.0, None, None, None],
        ["10001", "CTG02", "SOJA", 28, 13.7, None, 0.8, None, None, None, None, None],
        ["10001", "CTG03", "SOJA", 32, 14.5, 14.0, 1.2, None, None, None, None, None],
        ["10002", "CTG04", "MAIZ", 25, 15.2, 14.9, 1.4, None, 3.5, 3.1, None, None],
        ["10003", "CTG05", "GIRASOL", 20, 10.5, None, None, None, None, None, 44.0, 45.0],
        ["10004", "", "SOJA", 10, 13.0, None, None, None, None, None, None, None],
        ["10004", "CTG06", "CEBADA", 10, 13.0, None, None, None, None, None, None, None],
        ["10001", "CTG01", "SOJA", 30, 14.2, None, None, None, None, None, None, None],
    ]
    for r, fila in enumerate(ejemplos, start=2):
        for c, v in enumerate(fila, start=1):
            ws.cell(row=r, column=c, value=v).font = BASE_FONT

    # ---------- Hojas de resultados ----------
    salidas = {
        "Resumen Contratos": ["Contrato", "Cant. CTG", "Toneladas", "Consolidacion",
                              "Bonificacion %", "Rebaja %", "Neto %", "Merma %", "Advertencias"],
        "Resultados CTG": ["Contrato", "CTG", "Producto", "Toneladas", "Norma", "Version",
                           "Bonificacion %", "Rebaja %", "Neto %", "Merma %",
                           "Fecha calculo", "Advertencias"],
        "Detalle Calidad": ["Contrato", "CTG", "Producto", "Rubro", "Unidad", "Planta",
                            "Camara", "Valor valido", "Origen", "Base", "Norma", "Version",
                            "Regla", "Tipo ajuste", "Ajuste %"],
        "Errores": ["Fila", "Contrato", "CTG", "Motivo del rechazo"],
    }
    for nombre, headers in salidas.items():
        ws = wb.create_sheet(nombre)
        write_headers(ws, headers)

    # ---------- Trazabilidad ----------
    ws = wb.create_sheet("Trazabilidad")
    write_headers(ws, [
        "FechaHora", "Contrato", "CTG", "Producto", "Toneladas", "Rubro",
        "Planta", "Camara", "ValorUtilizado", "Origen", "Norma", "Version",
        "Regla", "TipoAjuste", "Ajuste %", "BonifCTG %", "RebajaCTG %",
        "NetoCTG %", "MermaCTG %",
    ], [19, 10, 8, 10, 10, 22, 8, 8, 13, 9, 40, 8, 45, 12, 9, 10, 11, 9, 10])

    wb.save(target)
    print(f"Libro generado: {target}")


if __name__ == "__main__":
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else \
        Path(__file__).parent / "CalculadoraCalidadGranos.xlsx"
    build(destino)
