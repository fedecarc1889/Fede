"""Tests de importación Excel, validaciones por fila y exportación."""

import pytest
from openpyxl import Workbook, load_workbook

from grain_quality.exporter import export_results
from grain_quality.importer import import_excel, parse_number
from grain_quality.models import Source


HEADERS = [
    "Contrato", "CTG", "Producto", "Toneladas",
    "Humedad Planta", "Humedad Camara",
    "Materias extranas Planta", "Materias extranas Camara",
]


def _build_xlsx(path, rows, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_parse_number():
    assert parse_number("14,2") == pytest.approx(14.2)
    assert parse_number("14.2") == pytest.approx(14.2)
    assert parse_number("1.234,56") == pytest.approx(1234.56)
    assert parse_number(0) == 0.0
    assert parse_number("") is None
    assert parse_number(None) is None
    with pytest.raises(ValueError):
        parse_number("abc")


def test_importacion_completa(tmp_path):
    path = _build_xlsx(tmp_path / "in.xlsx", [
        ["10001", "CTG01", "SOJA", 30, 14.2, 13.8, 1.5, 1.1],
        ["10001", "CTG02", "SOJA", 28, 13.7, None, 0.8, None],
        ["10001", "CTG03", "SOJA", 32, 14.5, 14.0, 1.2, None],
    ])
    result = import_excel(path)
    assert not result.errors
    assert len(result.ctg_results) == 3
    assert len(result.contract_results) == 1

    ctg1 = result.ctg_results[0]
    humedad = next(r for r in ctg1.rubros if r.rubro == "humedad")
    assert humedad.selected_value == 13.8
    assert humedad.selected_source == Source.CAMARA
    me = next(r for r in ctg1.rubros if r.rubro == "materias_extranas")
    assert me.selected_value == 1.1

    # CTG02 solo tiene Planta
    ctg2 = result.ctg_results[1]
    humedad2 = next(r for r in ctg2.rubros if r.rubro == "humedad")
    assert humedad2.selected_source == Source.PLANTA
    assert humedad2.selected_value == 13.7


def test_filas_con_error_no_frenan_el_archivo(tmp_path):
    path = _build_xlsx(tmp_path / "in.xlsx", [
        ["10001", "CTG01", "SOJA", 30, 14.0, None, None, None],
        ["10001", "", "SOJA", 30, 14.0, None, None, None],          # CTG vacío
        ["10001", "CTG02", "CEBADA", 30, 14.0, None, None, None],   # producto inexistente
        ["10001", "CTG01", "SOJA", 30, 14.0, None, None, None],     # CTG duplicado
        ["10001", "CTG03", "SOJA", 30, "abc", None, None, None],    # valor no numérico
        ["10001", "CTG04", "SOJA", 30, 13.0, None, None, None],
    ])
    result = import_excel(path)
    assert len(result.ctg_results) == 2  # CTG01 y CTG04
    motivos = " | ".join(e.reason for e in result.errors)
    assert len(result.errors) == 4
    assert "CTG vacío" in motivos
    assert "inexistente" in motivos
    assert "duplicado" in motivos
    assert "no numérico" in motivos
    # cada error identifica su fila
    assert all(e.row_number is not None for e in result.errors)


def test_columnas_obligatorias_faltantes(tmp_path):
    path = _build_xlsx(tmp_path / "in.xlsx", [], headers=["Contrato", "Toneladas"])
    result = import_excel(path)
    assert result.errors
    assert "obligatorias" in result.errors[0].reason


def test_celda_vacia_no_es_cero(tmp_path):
    path = _build_xlsx(tmp_path / "in.xlsx", [
        ["10001", "CTG01", "SOJA", 30, 14.0, None, None, None],
    ])
    result = import_excel(path)
    me = next(r for r in result.ctg_results[0].rubros
              if r.rubro == "materias_extranas")
    assert me.selected_value is None
    assert me.plant_value is None  # vacío != 0


def test_kilos_se_convierten_a_toneladas(tmp_path):
    headers = ["Contrato", "CTG", "Producto", "Kilos", "Humedad Planta"]
    path = _build_xlsx(tmp_path / "in.xlsx",
                       [["10001", "CTG01", "SOJA", 30000, 14.0]], headers)
    result = import_excel(path)
    assert result.ctg_results[0].weight_tn == pytest.approx(30.0)


def test_exportacion_cuatro_hojas(tmp_path):
    entrada = _build_xlsx(tmp_path / "in.xlsx", [
        ["10001", "CTG01", "SOJA", 30, 14.2, 13.8, 1.5, 1.1],
        ["10001", "", "SOJA", 30, 14.0, None, None, None],
    ])
    result = import_excel(entrada)
    salida = tmp_path / "out.xlsx"
    export_results(result.contract_results, result.errors, salida)

    workbook = load_workbook(salida)
    assert workbook.sheetnames == [
        "Resumen Contratos", "Resultados CTG", "Detalle Calidad", "Errores",
    ]
    resumen = list(workbook["Resumen Contratos"].values)
    assert len(resumen) == 2  # encabezado + 1 contrato
    detalle = list(workbook["Detalle Calidad"].values)
    assert len(detalle) > 1
    errores = list(workbook["Errores"].values)
    assert len(errores) == 2  # encabezado + 1 fila rechazada
